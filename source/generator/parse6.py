import re,json,datetime
from openpyxl import load_workbook
wf=load_workbook('../sq/w6.xlsx');wv=load_workbook('../sq/w6.xlsx',data_only=True)
def S(v):
    if v is None:return None
    if isinstance(v,float) and v.is_integer():return str(int(v))
    return str(v).strip()
MAXR={'B12':'BP','B13':'SQ','B14':'DL'}
def wexpr(f,prevmap):
    if f is None:return None
    if not isinstance(f,str):return {'t':'num','v':f}
    if not f.startswith('='):return {'t':'text','v':f.strip()}
    g=f.replace('$','').replace(' ','')
    m=re.search(r'IF\(Inputs!B\d+="kg",MROUND\(Inputs!(B\d+)\*([\d.]+)(?:\+([\d.]+))?,2\.5\)(-[\d.]+)?,MROUND\(Inputs!(B\d+)\*([\d.]+)',g)
    if m:
        kgref,fac,add,minus,lbref=m.group(1),float(m.group(2)),float(m.group(3) or 0),float(m.group(4) or 0),m.group(5)
        ref=MAXR[lbref]  # lb-Zweig ist korrekt (kg-Zweig verweist bei Backoff-Squats fälschlich auf das Kreuzheben-1RM)
        return {'t':'ref','ref':ref,'pct':round(fac*100,3),'add':add+minus,'bug':kgref!=lbref}
    m=re.fullmatch(r'=([A-Z])(\d+)\*([\d.]+)',g)
    if m:return {'t':'rel','cell':m.group(1)+m.group(2),'f':float(m.group(3))}
    m=re.fullmatch(r'=([A-Z])(\d+)',g)
    if m:return {'t':'eq','cell':m.group(1)+m.group(2)}
    return {'t':'?','f':f}
WD={0:'Montag',1:'Dienstag',3:'Donnerstag',4:'Freitag',5:'Samstag'}
weeks=[]
for wi,name in enumerate(['Week 1','Week 2','Week 3','Week 4','Week 5','Week 6']):
    ws=wf[name];v=wv[name]
    title=S(v['A1'].value)
    days=[];cur=None;wknotes=[]
    for r in range(3,ws.max_row+1):
        a=ws.cell(r,1).value;av=v.cell(r,1).value
        if isinstance(av,datetime.datetime) or (isinstance(a,str) and re.match(r'=(Inputs!B6|A\d+)',a.replace('$','')) and 'B17' not in a and 'B18' not in a):
            cur={'blocks':[],'notes':[]};days.append(cur);continue
        if cur is None or av in (None,''):continue
        bcol=S(v.cell(r,2).value)
        cells=[(ws.cell(r,c).value,S(v.cell(r,c+1).value)) for c in (3,5,7,9,11)]
        if bcol is None and all(x[0] is None and x[1] is None for x in cells) and not str(av).startswith(('Backoff',)):
            (cur['notes'] if len(days) else wknotes).append(S(av));continue
        if bcol and bcol.lower().startswith('no sets'):
            cur['notes'].append(S(av)+': '+bcol);continue
        sets=[]
        for c,(wfm,reps) in zip((3,5,7,9,11),cells):
            if wfm is None and reps is None:continue
            fl=ws.cell(r,c+1).fill
            if fl is not None and fl.fill_type=='solid' and str(fl.fgColor.rgb)=='FF000000':continue
            sets.append({'col':c,'w':wexpr(wfm,None),'reps':reps})
        cur['blocks'].append({'row':r,'name':S(av),'sets':sets})
    weeks.append({'title':title,'days':days,'notes':wknotes})
# Tage: Wochentage aus Inputs-Offsets (alle Wochen gleich): Mo, Di, Do, Fr, Sa
inputs={'info':[S(wv['Inputs'].cell(r,1).value) for r in range(21,26)]}
json.dump({'weeks':weeks,'inputs':inputs},open('parsed6.json','w'),ensure_ascii=False,indent=1)
for w in weeks:
    print(w['title'],[len(d['blocks']) for d in w['days']])
    for d in w['days']:
        for b in d['blocks']:
            u=[s['w'] for s in b['sets'] if s['w'] and s['w']['t'] in ('?',)]
            if u:print('  ?',b['name'],u)
