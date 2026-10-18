import re,json,datetime,sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
SRC='../sq/'
WD={'monday':0,'tuesday':1,'wednesday':2,'thursday':3,'friday':4,'saturday':5,'sunday':6}
def sval(v):
    if isinstance(v,datetime.datetime): return f'{v.month}-{v.day}'
    if isinstance(v,float) and v.is_integer(): return str(int(v))
    return None if v is None else str(v).strip()
def parse_book(fam,fname):
    wf=load_workbook(SRC+fname); wv=load_workbook(SRC+fname,data_only=True)
    out=[]
    for ws in wf.worksheets:
        if 'Mission' in ws.title: continue
        v=wv[ws.title]
        # week blocks from row 8
        starts=[(c.column,str(v.cell(8,c.column).value).strip()) for c in ws[8] if v.cell(8,c.column).value and 'WEEK' in str(v.cell(8,c.column).value).upper()]
        hasNotes=lambda col: str(v.cell(10,col+6).value or '').lower().find('note')>=0
        cellmap={}  # (col,row)->(w,d,b)
        weeks=[]
        maxr=max(r for r in range(1,ws.max_row+1) if any(ws.cell(r,c).value not in (None,'') for c in range(1,ws.max_column+1)))
        keyrow=next((r for r in range(10,maxr+1) if str(v.cell(r,2).value or '').strip().upper()=='KEY'),maxr+1)
        for wi,(col,label) in enumerate(starts):
            days=[];cur=None
            for r in range(10,keyrow):
                nm=v.cell(r,col).value; nmf=ws.cell(r,col).value
                s1=v.cell(r,col+1).value
                if isinstance(s1,str) and s1.strip().lower()=='sets':
                    t=str(nm or '').strip()
                    m=re.search(r'(monday|tuesday|wednesday|thursday|friday|saturday|sunday)',t,re.I)
                    cur={'name':t,'wd':WD[m.group(1).lower()] if m else None,'opt':'optional' in t.lower(),'blocks':[]};days.append(cur);continue
                if cur is None: continue
                if isinstance(s1,(int,float)) and nm not in (None,''):
                    b=len(cur['blocks'])
                    wcell=ws.cell(r,col+5).value
                    blk={'row':r,'name':str(nm).strip(),'sets':int(s1),'reps':sval(v.cell(r,col+2).value),'rpe':sval(v.cell(r,col+3).value),
                         'hype':sval(v.cell(r,col+4).value),'wf':wcell if isinstance(wcell,str) else sval(wcell),
                         'note':sval(v.cell(r,col+6).value) if hasNotes(col) else None,'col':col,'nmf':isinstance(nmf,str) and nmf.startswith('=')}
                    cur['blocks'].append(blk);cellmap[(col+5,r)]=(wi,len(days)-1,b)
                elif nm not in (None,'') and s1 in (None,''):
                    # text row (e.g. "Hack Squat Slot Dropped ...")
                    t=str(nm).strip()
                    if t.upper()=='OR' and cur['blocks']: cur['blocks'][-1]['orNext']=1
                    if len(t)>3: cur.setdefault('extra',[]).append(t)
            weeks.append({'label':label,'days':days})
        # resolve weight formulas
        for wi,wk in enumerate(weeks):
            for di,dy in enumerate(wk['days']):
                for bi,b in enumerate(dy['blocks']):
                    f=b.pop('wf');b['w']=None
                    if isinstance(f,str) and f.startswith('='):
                        m=re.fullmatch(r'=(?:MROUND\()?\$?([A-Z]{1,2})\$?(\d+)(?:\*([\d.]+))?(?:,[\d.]+\))?',f.replace(' ',''))
                        if not m: b['w']={'t':'?','f':f};continue
                        c,r,fac=CI(m.group(1)),int(m.group(2)),float(m.group(3) or 1)
                        if c==1: b['w']={'t':'ref','ref':'SQ','pct':round(fac*100,3)}
                        elif (c,r) in cellmap: b['w']={'t':'src','src':cellmap[(c,r)],'pct':round(fac*100,3)}
                        else: b['w']={'t':'?','f':f}
                    elif f: b['wtext']=f
                    b.pop('col');b.pop('row')
        key=[]
        for r in range(keyrow+1,maxr+1):
            a,bb=v.cell(r,2).value,v.cell(r,3).value
            if a or bb: key.append([sval(a),sval(bb)])
        out.append({'sheet':ws.title,'weeks':weeks,'key':key})
    return out
res={}
for fam,fn in [('pv','pv.xlsx'),('rm5','rm5.xlsx'),('rm20','rm20.xlsx'),('pk8','pk8.xlsx')]:
    res[fam]=parse_book(fam,fn)
json.dump(res,open('parsed.json','w'),ensure_ascii=False,indent=1)
for fam,vs in res.items():
    for x in vs:
        print(fam,x['sheet'][:28],[ (wk['label'][:18],[ (d['name'][:12],len(d['blocks'])) for d in wk['days']]) for wk in x['weeks']][:2])
        unk=[b for wk in x['weeks'] for d in wk['days'] for b in d['blocks'] if b['w'] and b['w']['t']=='?']
        if unk: print('  UNRESOLVED',[b['w'] for b in unk][:5])
